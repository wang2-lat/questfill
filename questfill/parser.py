"""问卷解析模块"""
from pathlib import Path
from typing import List, Dict
import pypdf
from docx import Document
from openpyxl import load_workbook

def parse_questionnaire(file_path: Path) -> List[Dict[str, str]]:
    """解析问卷文件，提取问题列表"""
    suffix = file_path.suffix.lower()
    
    if suffix == '.pdf':
        return parse_pdf(file_path)
    elif suffix in ['.docx', '.doc']:
        return parse_word(file_path)
    elif suffix in ['.xlsx', '.xls']:
        return parse_excel(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {suffix}")

def parse_pdf(file_path: Path) -> List[Dict[str, str]]:
    """解析 PDF 问卷"""
    questions = []
    reader = pypdf.PdfReader(str(file_path))
    
    # 提取表单字段
    if reader.get_fields():
        for field_name, field_data in reader.get_fields().items():
            questions.append({
                'id': field_name,
                'question': field_name.replace('_', ' ').title(),
                'type': 'text',
                'answer': ''
            })
    else:
        # 如果没有表单字段，提取文本中的问题（简单启发式）
        for page in reader.pages:
            text = page.extract_text()
            for line in text.split('\n'):
                if '?' in line or '：' in line or ':' in line:
                    questions.append({
                        'id': f"q_{len(questions)}",
                        'question': line.strip(),
                        'type': 'text',
                        'answer': ''
                    })
    
    return questions

def parse_word(file_path: Path) -> List[Dict[str, str]]:
    """解析 Word 问卷"""
    questions = []
    doc = Document(str(file_path))
    
    # 提取表格中的问题
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) >= 2 and cells[0]:
                questions.append({
                    'id': f"q_{len(questions)}",
                    'question': cells[0],
                    'type': 'text',
                    'answer': cells[1] if len(cells) > 1 else ''
                })
    
    # 提取段落中的问题
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and ('?' in text or '：' in text or ':' in text):
            questions.append({
                'id': f"q_{len(questions)}",
                'question': text,
                'type': 'text',
                'answer': ''
            })
    
    return questions

def parse_excel(file_path: Path) -> List[Dict[str, str]]:
    """解析 Excel 问卷"""
    questions = []
    wb = load_workbook(str(file_path))
    ws = wb.active
    
    # 假设第一列是问题，第二列是答案
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            questions.append({
                'id': f"q_{len(questions)}",
                'question': str(row[0]).strip(),
                'type': 'text',
                'answer': str(row[1]).strip() if len(row) > 1 and row[1] else ''
            })
    
    return questions