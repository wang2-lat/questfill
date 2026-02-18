"""文件导出模块"""
from pathlib import Path
from typing import Dict
import shutil
import pypdf
from docx import Document
from openpyxl import load_workbook

def export_questionnaire(input_file: Path, filled_data: Dict[str, str], output_file: Path):
    """导出填写完成的问卷"""
    suffix = input_file.suffix.lower()
    
    if suffix == '.pdf':
        export_pdf(input_file, filled_data, output_file)
    elif suffix in ['.docx', '.doc']:
        export_word(input_file, filled_data, output_file)
    elif suffix in ['.xlsx', '.xls']:
        export_excel(input_file, filled_data, output_file)

def export_pdf(input_file: Path, filled_data: Dict[str, str], output_file: Path):
    """导出填写后的 PDF"""
    reader = pypdf.PdfReader(str(input_file))
    writer = pypdf.PdfWriter()
    
    # 复制所有页面
    for page in reader.pages:
        writer.add_page(page)
    
    # 填充表单字段
    if reader.get_fields():
        writer.update_page_form_field_values(
            writer.pages[0], filled_data
        )
    
    with open(output_file, 'wb') as f:
        writer.write(f)

def export_word(input_file: Path, filled_data: Dict[str, str], output_file: Path):
    """导出填写后的 Word"""
    doc = Document(str(input_file))
    
    # 填充表格
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            q_id = f"q_{i}"
            if q_id in filled_data and len(row.cells) >= 2:
                row.cells[1].text = filled_data[q_id]
    
    doc.save(str(output_file))

def export_excel(input_file: Path, filled_data: Dict[str, str], output_file: Path):
    """导出填写后的 Excel"""
    wb = load_workbook(str(input_file))
    ws = wb.active
    
    # 填充答案列（假设第二列）
    for i, row in enumerate(ws.iter_rows(min_row=2), start=0):
        q_id = f"q_{i}"
        if q_id in filled_data and len(row) >= 2:
            row[1].value = filled_data[q_id]
    
    wb.save(str(output_file))